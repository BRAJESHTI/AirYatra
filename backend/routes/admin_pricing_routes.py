"""Admin/CEO: fixed route pricing + custom quotes on operator's behalf + AirYatra own fleet"""
from fastapi import APIRouter, HTTPException, Depends
from pydantic import BaseModel
from typing import Optional
from datetime import datetime, timezone, timedelta
import uuid

from database import get_database
from middleware import get_current_user

router = APIRouter(prefix="/admin/pricing", tags=["Admin Pricing"])

ADMIN_ROLES = {"admin", "super_admin", "ceo"}
OWN_FLEET_OPERATOR_ID = "airyatra_own_fleet"
OWN_FLEET_NAME = "AirYatra Own Fleet"


def _require_admin(user: dict):
    if not ADMIN_ROLES & set(user.get("roles", [])):
        raise HTTPException(status_code=403, detail="Admin/CEO access required")


# ==================== HELPERS FOR ADMIN UI ====================

@router.get("/operators")
async def list_operators(user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    ops = await db.operators.find(
        {}, {"_id": 0, "id": 1, "company_name": 1, "status": 1}).to_list(200)
    ops.insert(0, {"id": OWN_FLEET_OPERATOR_ID, "company_name": OWN_FLEET_NAME, "status": "active"})
    return {"operators": ops}


@router.get("/aircraft")
async def list_operator_aircraft(operator_id: str, user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    aircraft = await db.aircraft.find(
        {"operator_id": operator_id},
        {"_id": 0, "id": 1, "model_name": 1, "name": 1, "registration_number": 1,
         "service_category": 1, "capacity": 1, "hourly_rate": 1}).to_list(100)
    return {"aircraft": aircraft}


# ==================== FIXED ROUTE ON OPERATOR'S BEHALF ====================

class AdminFixedRouteCreate(BaseModel):
    operator_id: str
    origin: str
    destination: str
    aircraft_type: str = "helicopter"
    aircraft_id: Optional[str] = None
    base_price: float
    price_per_seat: Optional[float] = None
    route_name: Optional[str] = None


@router.post("/fixed-route")
async def create_fixed_route_on_behalf(body: AdminFixedRouteCreate, user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    if body.base_price <= 0:
        raise HTTPException(status_code=400, detail="base_price must be positive")

    if body.operator_id == OWN_FLEET_OPERATOR_ID:
        operator_name = OWN_FLEET_NAME
    else:
        operator = await db.operators.find_one({"id": body.operator_id}, {"_id": 0, "company_name": 1})
        if not operator:
            raise HTTPException(status_code=404, detail="Operator not found")
        operator_name = operator["company_name"]

    origin, destination = body.origin.strip().title(), body.destination.strip().title()
    existing = await db.fixed_routes.find_one({
        "operator_id": body.operator_id, "origin": origin, "destination": destination,
        "aircraft_type": body.aircraft_type, "status": "active", "is_deleted": {"$ne": True}})
    if existing:
        raise HTTPException(status_code=400, detail="Active route already exists for this operator")

    aircraft = None
    if body.aircraft_id:
        aircraft = await db.aircraft.find_one({"id": body.aircraft_id}, {"_id": 0})

    route_doc = {
        "id": str(uuid.uuid4()),
        "route_code": f"AYR-{str(uuid.uuid4())[:6].upper()}",
        "route_name": body.route_name or f"{origin} to {destination} Express",
        "operator_id": body.operator_id,
        "operator_name": operator_name,
        "origin": origin,
        "destination": destination,
        "aircraft_type": body.aircraft_type,
        "aircraft_id": body.aircraft_id,
        "aircraft_name": (aircraft or {}).get("model_name"),
        "base_price": body.base_price,
        "price_per_seat": body.price_per_seat or 0,
        "demand_multiplier": 1.0,
        "status": "active",
        "is_active": True,
        "is_deleted": False,
        "added_by_admin": True,
        "added_by_name": user.get("full_name") or user.get("email"),
        "added_by_role": "ceo" if "ceo" in user.get("roles", []) else "admin",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.fixed_routes.insert_one({**route_doc})
    return {"message": "Fixed route added on operator's behalf", "route": route_doc}


@router.get("/fixed-routes")
async def list_fixed_routes(operator_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    query = {"is_deleted": {"$ne": True}}
    if operator_id:
        query["operator_id"] = operator_id
    routes = await db.fixed_routes.find(query, {"_id": 0}).sort("created_at", -1).to_list(300)
    return {"routes": routes, "total": len(routes)}


@router.delete("/fixed-route/{route_id}")
async def deactivate_fixed_route(route_id: str, user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    result = await db.fixed_routes.update_one(
        {"id": route_id},
        {"$set": {"status": "inactive", "is_active": False,
                  "deactivated_by": user["id"], "deactivated_at": datetime.now(timezone.utc).isoformat()}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Route not found")
    return {"message": "Route deactivated"}


# ==================== QUOTE ON OPERATOR'S BEHALF ====================

@router.get("/pending-quote-bookings")
async def pending_quote_bookings(user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    query = {"status": {"$in": ["pending", "quote_requested", "quotes_received", "broadcasted"]}}
    proj = {"_id": 0, "id": 1, "booking_number": 1, "inquiry_number": 1, "from_location": 1,
            "to_location": 1, "departure_date": 1, "travel_date": 1, "departure_time": 1,
            "customer_name": 1, "status": 1, "total_passengers": 1, "passengers": 1}
    bookings = await db.bookings.find(query, proj).sort("created_at", -1).to_list(20)
    inquiries = await db.inquiries.find(query, proj).sort("created_at", -1).to_list(20)
    seen, merged = set(), []
    for b in bookings + inquiries:
        if b["id"] not in seen:
            seen.add(b["id"])
            merged.append(b)
    return {"bookings": merged[:30]}


class AdminQuoteCreate(BaseModel):
    booking_id: str
    operator_id: str
    amount: float
    notes: Optional[str] = None


@router.post("/quote-on-behalf")
async def submit_quote_on_behalf(body: AdminQuoteCreate, user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    if body.amount <= 0:
        raise HTTPException(status_code=400, detail="Valid quote amount required")

    booking = await db.bookings.find_one({"id": body.booking_id}, {"_id": 0}) or \
              await db.inquiries.find_one({"id": body.booking_id}, {"_id": 0})
    if not booking:
        raise HTTPException(status_code=404, detail="Booking not found")

    if body.operator_id == OWN_FLEET_OPERATOR_ID:
        operator_name = OWN_FLEET_NAME
    else:
        operator = await db.operators.find_one({"id": body.operator_id}, {"_id": 0, "company_name": 1})
        if not operator:
            raise HTTPException(status_code=404, detail="Operator not found")
        operator_name = operator["company_name"]

    from services.platform_fee_service import resolve_platform_fee, compute_platform_fee
    from services.dynamic_pricing_service import get_urgency_settings, get_urgency_percent
    fee_rule = await resolve_platform_fee(
        db, booking.get("from_location") or booking.get("pickup_location", ""),
        booking.get("to_location") or booking.get("drop_location", ""))
    platform_fee = compute_platform_fee(body.amount, fee_rule)
    urgency_pct, urgency_label = get_urgency_percent(
        await get_urgency_settings(db),
        booking.get("departure_date") or booking.get("travel_date") or "",
        booking.get("departure_time") or booking.get("travel_time"))
    urgency_surcharge = round(body.amount * urgency_pct / 100, 2)
    customer_total = round(body.amount + platform_fee + urgency_surcharge, 2)

    quote_id = str(uuid.uuid4())
    quote = {
        "id": quote_id,
        "booking_id": body.booking_id,
        "operator_id": body.operator_id,
        "operator_name": operator_name,
        "customer_id": booking.get("customer_id") or booking.get("user_id"),
        "amount": customer_total,
        "quoted_price": customer_total,
        "operator_payout": round(body.amount, 2),
        "platform_fee": platform_fee,
        "platform_fee_rule": fee_rule.get("label"),
        "urgency_percent": urgency_pct,
        "urgency_surcharge": urgency_surcharge,
        "urgency_label": urgency_label,
        "notes": body.notes,
        "status": "pending",
        "submitted_by_admin": True,
        "submitted_by_name": user.get("full_name") or user.get("email"),
        "submitted_by_role": "ceo" if "ceo" in user.get("roles", []) else "admin",
        "validity_hours": 24,
        "valid_until": (datetime.now(timezone.utc) + timedelta(hours=24)).isoformat(),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.quotes.insert_one({**quote})

    for coll in (db.bookings, db.inquiries):
        await coll.update_one(
            {"id": body.booking_id},
            {"$set": {"status": "quotes_received", "updated_at": datetime.now(timezone.utc).isoformat()}})

    customer_user_id = quote["customer_id"]
    if customer_user_id:
        await db.in_app_notifications.insert_one({
            "id": str(uuid.uuid4()),
            "user_id": customer_user_id,
            "type": "quote_received",
            "title": f"✈️ New Quote Received - ₹{customer_total:,.0f}",
            "message": f"{operator_name} ne aapki booking ke liye quote bheja hai. Abhi review karein!",
            "reference_id": body.booking_id,
            "data": {"quote_id": quote_id, "amount": customer_total, "operator_name": operator_name},
            "read": False,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    return {"message": "Quote submitted on operator's behalf", "quote": quote}


# ==================== AIRYATRA OWN FLEET ====================

class OwnAircraftCreate(BaseModel):
    model_name: str
    service_category: str = "helicopter"
    registration_number: str
    capacity: int
    hourly_rate: float
    engine_type: Optional[str] = None
    engine_model: Optional[str] = None
    base_city: str = "Mumbai"
    base_lat: Optional[float] = None
    base_lng: Optional[float] = None
    pilot_experience_hours: int = 3000
    ownership: str = "own"


class OwnAircraftUpdate(BaseModel):
    hourly_rate: Optional[float] = None
    is_available: Optional[bool] = None
    capacity: Optional[int] = None
    engine_type: Optional[str] = None
    engine_model: Optional[str] = None
    pilot_experience_hours: Optional[int] = None


CATEGORY_IMAGES = {
    "helicopter": "/services/helicopter.jpg",
    "chartered_plane": "/services/private_jet.jpg",
    "air_ambulance": "/services/air_ambulance.jpg",
    "yacht_cruiser": "/services/yacht_cruiser.jpg",
    "cargo": "/services/cargo.jpg",
    "joy_ride": "/services/joy_ride.jpg",
}


@router.get("/own-aircraft")
async def list_own_aircraft(user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    aircraft = await db.aircraft.find(
        {"operator_id": OWN_FLEET_OPERATOR_ID}, {"_id": 0}).sort("created_at", -1).to_list(100)
    return {"aircraft": aircraft, "total": len(aircraft)}


@router.post("/own-aircraft")
async def add_own_aircraft(body: OwnAircraftCreate, user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    if body.hourly_rate <= 0 or body.capacity <= 0:
        raise HTTPException(status_code=400, detail="hourly_rate and capacity must be positive")
    existing = await db.aircraft.find_one({"registration_number": body.registration_number}, {"_id": 0, "id": 1})
    if existing:
        raise HTTPException(status_code=400, detail="Aircraft with this registration already exists")

    doc = {
        "id": f"own-{str(uuid.uuid4())[:8]}",
        "operator_id": OWN_FLEET_OPERATOR_ID,
        "operator_name": OWN_FLEET_NAME,
        "own_fleet": True,
        "ownership": body.ownership,
        "model_name": body.model_name,
        "aircraft_type": body.model_name,
        "service_category": body.service_category,
        "registration_number": body.registration_number,
        "capacity": body.capacity,
        "hourly_rate": body.hourly_rate,
        "engine_type": body.engine_type,
        "engine_model": body.engine_model,
        "base_location": body.base_city,
        "base_coordinates": {"city": body.base_city, "lat": body.base_lat, "lng": body.base_lng},
        "pilot_experience_hours": body.pilot_experience_hours,
        "cabin_crew": 0,
        "amenities": {},
        "image": CATEGORY_IMAGES.get(body.service_category, "/services/helicopter.jpg"),
        "marketplace_listed": True,
        "is_available": True,
        "verified": True,
        "rating": 4.8,
        "total_flights": 0,
        "added_by": user["id"],
        "added_by_name": user.get("full_name") or user.get("email"),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.aircraft.insert_one({**doc})
    return {"message": "AirYatra own aircraft added to marketplace", "aircraft": doc}


@router.put("/own-aircraft/{aircraft_id}")
async def update_own_aircraft(aircraft_id: str, body: OwnAircraftUpdate, user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    update = {k: v for k, v in body.dict().items() if v is not None}
    if not update:
        raise HTTPException(status_code=400, detail="Nothing to update")
    if "hourly_rate" in update and update["hourly_rate"] <= 0:
        raise HTTPException(status_code=400, detail="hourly_rate must be positive")
    update["updated_at"] = datetime.now(timezone.utc).isoformat()
    result = await db.aircraft.update_one(
        {"id": aircraft_id, "operator_id": OWN_FLEET_OPERATOR_ID}, {"$set": update})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Own aircraft not found")
    return {"message": "Aircraft updated"}


@router.delete("/own-aircraft/{aircraft_id}")
async def remove_own_aircraft(aircraft_id: str, user: dict = Depends(get_current_user)):
    _require_admin(user)
    db = get_database()
    result = await db.aircraft.delete_one({"id": aircraft_id, "operator_id": OWN_FLEET_OPERATOR_ID})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Own aircraft not found")
    return {"message": "Aircraft removed from marketplace"}


# ==================== REPORTS ====================

PAID_STATUSES = {"paid", "fully_paid", "captured"}


def _created_range(start_date: Optional[str], end_date: Optional[str]) -> Optional[dict]:
    rng = {}
    if start_date:
        rng["$gte"] = f"{start_date[:10]}T00:00:00"
    if end_date:
        rng["$lte"] = f"{end_date[:10]}T23:59:59.999999+00:00"
    return rng or None


@router.get("/own-fleet-bookings")
async def own_fleet_bookings(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    user: dict = Depends(get_current_user),
):
    """Bookings landed on AirYatra's own aircraft + earnings"""
    _require_admin(user)
    db = get_database()
    query = {"$or": [{"operator_id": OWN_FLEET_OPERATOR_ID}, {"aircraft_id": {"$regex": "^own-"}}]}
    rng = _created_range(start_date, end_date)
    if rng:
        query["created_at"] = rng
    proj = {"_id": 0, "id": 1, "inquiry_number": 1, "booking_number": 1, "from_location": 1,
            "to_location": 1, "travel_date": 1, "departure_date": 1, "customer_name": 1,
            "aircraft_model": 1, "aircraft_id": 1, "final_price": 1, "total_amount": 1,
            "estimated_price": 1, "payment_status": 1, "status": 1, "created_at": 1}
    rows = await db.inquiries.find(query, proj).sort("created_at", -1).to_list(200)
    rows += await db.bookings.find(query, proj).sort("created_at", -1).to_list(200)
    seen, bookings = set(), []
    for b in rows:
        if b["id"] not in seen:
            seen.add(b["id"])
            b["amount"] = b.get("final_price") or b.get("total_amount") or b.get("estimated_price") or 0
            b["is_paid"] = (b.get("payment_status") or "") in PAID_STATUSES
            bookings.append(b)
    bookings.sort(key=lambda b: b.get("created_at") or "", reverse=True)
    paid = [b for b in bookings if b["is_paid"]]
    return {
        "bookings": bookings,
        "summary": {
            "total_bookings": len(bookings),
            "paid_bookings": len(paid),
            "gross_earnings": round(sum(b["amount"] for b in paid), 2),
            "pending_amount": round(sum(b["amount"] for b in bookings if not b["is_paid"]), 2),
        },
    }


@router.get("/fee-revenue-report")
async def fee_revenue_report(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    user: dict = Depends(get_current_user),
):
    """Platform fee + urgency + surge income grouped by route and city"""
    _require_admin(user)
    db = get_database()
    rng = _created_range(start_date, end_date)

    routes: dict = {}

    def bucket(from_loc, to_loc):
        f = (from_loc or "Unknown").split(",")[0].strip().title()
        t = (to_loc or "Unknown").split(",")[0].strip().title()
        key = f"{f} → {t}"
        if key not in routes:
            routes[key] = {"route": key, "from_city": f, "to_city": t, "quotes": 0, "paid_bookings": 0,
                           "platform_fee": 0.0, "urgency_income": 0.0, "surge_income": 0.0,
                           "convenience_fee": 0.0, "realized_income": 0.0}
        return routes[key]

    # 1) Operator/admin quotes (custom quote flow)
    quotes_query = {"$or": [{"platform_fee": {"$gt": 0}}, {"urgency_surcharge": {"$gt": 0}}]}
    if rng:
        quotes_query["created_at"] = rng
    quotes = await db.quotes.find(
        quotes_query,
        {"_id": 0, "booking_id": 1, "platform_fee": 1, "urgency_surcharge": 1, "status": 1}
    ).sort("created_at", -1).to_list(500)
    booking_ids = list({q["booking_id"] for q in quotes if q.get("booking_id")})
    route_map = {}
    for coll in (db.bookings, db.inquiries):
        async for b in coll.find({"id": {"$in": booking_ids}},
                                 {"_id": 0, "id": 1, "from_location": 1, "to_location": 1, "payment_status": 1}):
            route_map.setdefault(b["id"], b)
    for q in quotes:
        b = route_map.get(q.get("booking_id")) or {}
        r = bucket(b.get("from_location"), b.get("to_location"))
        r["quotes"] += 1
        fee = float(q.get("platform_fee") or 0)
        urg = float(q.get("urgency_surcharge") or 0)
        r["platform_fee"] += fee
        r["urgency_income"] += urg
        if q.get("status") == "accepted" or (b.get("payment_status") or "") in PAID_STATUSES:
            r["realized_income"] += fee + urg

    # 2) Marketplace instant bookings (convenience fee + surge + urgency), paid only
    mkt_query = {"pricing_breakdown": {"$exists": True}, "payment_status": {"$in": list(PAID_STATUSES)}}
    if rng:
        mkt_query["created_at"] = rng
    async for b in db.inquiries.find(
            mkt_query,
            {"_id": 0, "from_location": 1, "to_location": 1, "pricing_breakdown": 1}).limit(500):
        p = b.get("pricing_breakdown") or {}
        r = bucket(b.get("from_location"), b.get("to_location"))
        r["paid_bookings"] += 1
        conv = float(p.get("convenience_fee") or 0)
        surge = float(p.get("surge_amount") or 0)
        urg = float(p.get("urgency_amount") or 0)
        r["convenience_fee"] += conv
        r["surge_income"] += surge
        r["urgency_income"] += urg
        r["realized_income"] += conv + surge + urg

    route_list = sorted(routes.values(), key=lambda r: -(r["platform_fee"] + r["surge_income"] + r["urgency_income"] + r["convenience_fee"]))
    for r in route_list:
        for k in ("platform_fee", "urgency_income", "surge_income", "convenience_fee", "realized_income"):
            r[k] = round(r[k], 2)

    cities: dict = {}
    for r in route_list:
        c = cities.setdefault(r["from_city"], {"city": r["from_city"], "routes": 0, "platform_fee": 0.0,
                                               "urgency_income": 0.0, "surge_income": 0.0,
                                               "convenience_fee": 0.0, "realized_income": 0.0})
        c["routes"] += 1
        for k in ("platform_fee", "urgency_income", "surge_income", "convenience_fee", "realized_income"):
            c[k] = round(c[k] + r[k], 2)
    city_list = sorted(cities.values(), key=lambda c: -c["realized_income"])

    totals = {k: round(sum(r[k] for r in route_list), 2)
              for k in ("platform_fee", "urgency_income", "surge_income", "convenience_fee", "realized_income")}
    totals["total_quotes"] = sum(r["quotes"] for r in route_list)
    totals["total_paid_bookings"] = sum(r["paid_bookings"] for r in route_list)

    return {"routes": route_list, "cities": city_list, "totals": totals}


@router.get("/revenue-trend")
async def revenue_trend(weeks: int = 8, user: dict = Depends(get_current_user)):
    """Weekly income trend: quote fees + marketplace fees + own-fleet paid earnings"""
    _require_admin(user)
    db = get_database()
    weeks = max(2, min(weeks, 26))
    now = datetime.now(timezone.utc)
    start_of_week = (now - timedelta(days=now.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    buckets = []
    for i in range(weeks - 1, -1, -1):
        ws = start_of_week - timedelta(weeks=i)
        buckets.append({
            "week_start": ws.strftime("%Y-%m-%d"),
            "label": ws.strftime("%d %b"),
            "fee_income": 0.0,
            "marketplace_income": 0.0,
            "own_fleet_earnings": 0.0,
            "total": 0.0,
        })
    range_start = buckets[0]["week_start"]

    def bucket_for(created_at):
        d = (created_at or "")[:10]
        for b in reversed(buckets):
            if d >= b["week_start"]:
                return b
        return None

    async for q in db.quotes.find(
            {"created_at": {"$gte": range_start},
             "$or": [{"platform_fee": {"$gt": 0}}, {"urgency_surcharge": {"$gt": 0}}]},
            {"_id": 0, "created_at": 1, "platform_fee": 1, "urgency_surcharge": 1}).limit(2000):
        b = bucket_for(q.get("created_at"))
        if b:
            b["fee_income"] += float(q.get("platform_fee") or 0) + float(q.get("urgency_surcharge") or 0)

    async for m in db.inquiries.find(
            {"created_at": {"$gte": range_start}, "pricing_breakdown": {"$exists": True},
             "payment_status": {"$in": list(PAID_STATUSES)}},
            {"_id": 0, "created_at": 1, "pricing_breakdown": 1, "operator_id": 1, "final_price": 1}).limit(2000):
        b = bucket_for(m.get("created_at"))
        if not b:
            continue
        p = m.get("pricing_breakdown") or {}
        b["marketplace_income"] += float(p.get("convenience_fee") or 0) + float(p.get("surge_amount") or 0) + float(p.get("urgency_amount") or 0)
        if m.get("operator_id") == OWN_FLEET_OPERATOR_ID:
            b["own_fleet_earnings"] += float(m.get("final_price") or 0)

    for b in buckets:
        for k in ("fee_income", "marketplace_income", "own_fleet_earnings"):
            b[k] = round(b[k], 2)
        b["total"] = round(b["fee_income"] + b["marketplace_income"] + b["own_fleet_earnings"], 2)
    return {"weeks": buckets}


@router.get("/invoice-email-log")
async def invoice_email_log(limit: int = 50, user: dict = Depends(get_current_user)):
    """Audit log of auto-sent invoice emails"""
    _require_admin(user)
    db = get_database()
    logs = await db.invoice_email_log.find({}, {"_id": 0}).sort("created_at", -1).to_list(min(limit, 200))
    sent = sum(1 for l in logs if l.get("status") == "sent")
    return {"logs": logs, "total": len(logs), "sent": sent}


@router.get("/operator-scorecards")
async def operator_scorecards(user: dict = Depends(get_current_user)):
    """Rank operators by quotes won, ratings and on-time flights"""
    _require_admin(user)
    db = get_database()

    operators = await db.operators.find({}, {"_id": 0, "id": 1, "company_name": 1, "status": 1, "rating": 1}).to_list(200)
    quotes = await db.quotes.find({}, {"_id": 0, "operator_id": 1, "status": 1}).to_list(2000)
    aircraft = await db.aircraft.find({}, {"_id": 0, "id": 1, "operator_id": 1, "rating": 1}).to_list(1000)
    records = await db.flight_records.find(
        {}, {"_id": 0, "aircraft_id": 1, "booking_id": 1, "departure_time": 1}).to_list(2000)

    ac_owner = {a["id"]: a.get("operator_id") for a in aircraft}
    booking_ids = list({r["booking_id"] for r in records if r.get("booking_id")})
    sched = {}
    for coll in (db.bookings, db.inquiries):
        async for b in coll.find({"id": {"$in": booking_ids}},
                                 {"_id": 0, "id": 1, "travel_date": 1, "departure_date": 1,
                                  "pickup_time": 1, "departure_time": 1, "travel_time": 1}):
            sched.setdefault(b["id"], b)

    def parse_dt(val):
        try:
            return datetime.fromisoformat(str(val).replace("Z", "+00:00"))
        except Exception:
            return None

    stats = {}
    for op in operators:
        stats[op["id"]] = {"operator_id": op["id"], "company_name": op["company_name"],
                           "status": op.get("status"), "quotes_sent": 0, "quotes_won": 0,
                           "flights_total": 0, "on_time_flights": 0, "on_time_known": 0,
                           "_ratings": []}

    for q in quotes:
        s = stats.get(q.get("operator_id"))
        if s:
            s["quotes_sent"] += 1
            if q.get("status") == "accepted":
                s["quotes_won"] += 1

    for a in aircraft:
        s = stats.get(a.get("operator_id"))
        if s and a.get("rating"):
            s["_ratings"].append(float(a["rating"]))

    for r in records:
        op_id = ac_owner.get(r.get("aircraft_id"))
        s = stats.get(op_id)
        if not s:
            continue
        s["flights_total"] += 1
        b = sched.get(r.get("booking_id")) or {}
        sched_date = (b.get("travel_date") or b.get("departure_date") or "")[:10]
        sched_time = (b.get("pickup_time") or b.get("departure_time") or b.get("travel_time") or "")[:5]
        actual = parse_dt(r.get("departure_time"))
        if sched_date and sched_time and actual:
            planned = parse_dt(f"{sched_date}T{sched_time}:00+00:00")
            if planned:
                s["on_time_known"] += 1
                delay_min = (actual.replace(tzinfo=None) - planned.replace(tzinfo=None)).total_seconds() / 60
                if delay_min <= 30:
                    s["on_time_flights"] += 1

    cards = []
    for op in operators:
        s = stats[op["id"]]
        win_rate = round(s["quotes_won"] / s["quotes_sent"] * 100, 1) if s["quotes_sent"] else 0.0
        rating = round(sum(s["_ratings"]) / len(s["_ratings"]), 2) if s["_ratings"] else float(op.get("rating") or 0) or None
        on_time = round(s["on_time_flights"] / s["on_time_known"] * 100, 1) if s["on_time_known"] else None

        participation = min(s["quotes_sent"] / 10, 1) * 100
        score = round(
            win_rate * 0.40 + ((rating or 0) / 5 * 100) * 0.30 + (on_time or 0) * 0.25 + participation * 0.05, 1)
        grade = "A+" if score >= 70 else "A" if score >= 55 else "B" if score >= 35 else "C" if score > 0 else "—"

        cards.append({"operator_id": s["operator_id"], "company_name": s["company_name"],
                      "status": s["status"], "quotes_sent": s["quotes_sent"], "quotes_won": s["quotes_won"],
                      "flights_total": s["flights_total"], "win_rate": win_rate, "rating": rating,
                      "on_time_percent": on_time, "score": score, "grade": grade})

    cards.sort(key=lambda c: (-c["score"], -c["quotes_won"], -(c["rating"] or 0)))
    for i, c in enumerate(cards):
        c["rank"] = i + 1
    return {"scorecards": cards, "total_operators": len(cards)}


@router.get("/export")
async def export_revenue_report(
    format: str = "excel",
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    user: dict = Depends(get_current_user),
):
    """Download revenue report as Excel or PDF"""
    _require_admin(user)
    fee = await fee_revenue_report(start_date=start_date, end_date=end_date, user=user)
    own = await own_fleet_bookings(start_date=start_date, end_date=end_date, user=user)
    period = f"{start_date or 'Beginning'} to {end_date or 'Today'}"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")

    import io
    from fastapi.responses import StreamingResponse

    if format == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title="AirYatra Revenue Report")
        styles = getSampleStyleSheet()
        story = [
            Paragraph("AirYatra - Revenue Report", styles["Title"]),
            Paragraph(f"Period: {period} | Generated: {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')}", styles["Normal"]),
            Spacer(1, 12),
            Paragraph("Summary", styles["Heading2"]),
        ]
        t = fee["totals"]
        s = own["summary"]
        summary_rows = [
            ["Platform Fees", f"Rs {t['platform_fee']:,.0f}", "Urgency Income", f"Rs {t['urgency_income']:,.0f}"],
            ["Surge Income", f"Rs {t['surge_income']:,.0f}", "Convenience Fees", f"Rs {t['convenience_fee']:,.0f}"],
            ["Realized Income", f"Rs {t['realized_income']:,.0f}", "Total Quotes", str(t["total_quotes"])],
            ["Own Fleet Bookings", str(s["total_bookings"]), "Own Fleet Gross Earnings", f"Rs {s['gross_earnings']:,.0f}"],
        ]
        st = Table(summary_rows, colWidths=[160, 120, 160, 120])
        st.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
            ("BACKGROUND", (2, 0), (2, -1), colors.whitesmoke),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story += [st, Spacer(1, 16), Paragraph("Route-wise Income", styles["Heading2"])]

        route_rows = [["Route", "Quotes", "Paid Bkgs", "Platform Fee", "Urgency", "Surge", "Convenience", "Realized"]]
        for r in fee["routes"][:40]:
            route_rows.append([r["route"], r["quotes"], r["paid_bookings"], f"{r['platform_fee']:,.0f}",
                               f"{r['urgency_income']:,.0f}", f"{r['surge_income']:,.0f}",
                               f"{r['convenience_fee']:,.0f}", f"{r['realized_income']:,.0f}"])
        rt = Table(route_rows, repeatRows=1)
        rt.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f97316")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ]))
        story += [rt, Spacer(1, 16), Paragraph("Own Fleet Bookings", styles["Heading2"])]

        own_rows = [["Ref", "Route", "Aircraft", "Date", "Customer", "Amount", "Status"]]
        for b in own["bookings"][:40]:
            own_rows.append([
                b.get("booking_number") or b.get("inquiry_number") or b["id"][:8],
                f"{b.get('from_location')} - {b.get('to_location')}",
                b.get("aircraft_model") or "", (b.get("travel_date") or b.get("departure_date") or "")[:10],
                b.get("customer_name") or "", f"{b['amount']:,.0f}",
                "Paid" if b["is_paid"] else (b.get("payment_status") or "pending"),
            ])
        ot = Table(own_rows, repeatRows=1)
        ot.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(ot)
        doc.build(story)
        pdf = buf.getvalue()
        return StreamingResponse(
            io.BytesIO(pdf), media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="AirYatra_Revenue_Report_{stamp}.pdf"',
                "Content-Length": str(len(pdf)),
                "Cache-Control": "no-cache",
            })

    # Excel
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="F97316")
    header_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "Summary"
    t, s = fee["totals"], own["summary"]
    ws.append(["AirYatra Revenue Report", period])
    ws.append([])
    for row in [("Platform Fees", t["platform_fee"]), ("Urgency Income", t["urgency_income"]),
                ("Surge Income", t["surge_income"]), ("Convenience Fees", t["convenience_fee"]),
                ("Realized Income", t["realized_income"]), ("Total Quotes", t["total_quotes"]),
                ("Own Fleet Bookings", s["total_bookings"]), ("Own Fleet Paid Bookings", s["paid_bookings"]),
                ("Own Fleet Gross Earnings", s["gross_earnings"]), ("Own Fleet Pending Amount", s["pending_amount"])]:
        ws.append(list(row))
    ws["A1"].font = Font(bold=True, size=14)

    ws2 = wb.create_sheet("Route Income")
    headers = ["Route", "Quotes", "Paid Bookings", "Platform Fee", "Urgency", "Surge", "Convenience", "Realized"]
    ws2.append(headers)
    for c in ws2[1]:
        c.fill, c.font = header_fill, header_font
    for r in fee["routes"]:
        ws2.append([r["route"], r["quotes"], r["paid_bookings"], r["platform_fee"],
                    r["urgency_income"], r["surge_income"], r["convenience_fee"], r["realized_income"]])

    ws3 = wb.create_sheet("Own Fleet Bookings")
    headers3 = ["Ref", "From", "To", "Aircraft", "Date", "Customer", "Amount", "Payment Status"]
    ws3.append(headers3)
    for c in ws3[1]:
        c.fill, c.font = header_fill, header_font
    for b in own["bookings"]:
        ws3.append([b.get("booking_number") or b.get("inquiry_number") or b["id"][:8],
                    b.get("from_location"), b.get("to_location"), b.get("aircraft_model"),
                    (b.get("travel_date") or b.get("departure_date") or "")[:10],
                    b.get("customer_name"), b["amount"],
                    "Paid" if b["is_paid"] else (b.get("payment_status") or "pending")])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            sheet.column_dimensions[col[0].column_letter].width = min(width + 3, 40)

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="AirYatra_Revenue_Report_{stamp}.xlsx"',
            "Content-Length": str(len(data)),
            "Cache-Control": "no-cache",
        })
