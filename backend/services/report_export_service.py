"""
AirYatra Report Export Service
Generate Excel reports for bookings, finance, and analytics
"""

import io
import logging
from datetime import datetime, timezone, timedelta
from typing import Dict, Any, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Styles
HEADER_FILL = PatternFill(start_color="FF6B00", end_color="FF6B00", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
SUBHEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
MONEY_FILL = PatternFill(start_color="e8f5e9", end_color="e8f5e9", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)


class ReportExportService:
    """Service for generating Excel reports"""
    
    def __init__(self):
        pass
    
    def generate_booking_report(
        self,
        bookings: List[Dict],
        title: str = "Booking Report",
        date_range: Optional[tuple] = None
    ) -> io.BytesIO:
        """
        Generate Excel report for bookings
        
        Args:
            bookings: List of booking records
            title: Report title
            date_range: Optional (start_date, end_date) tuple
        
        Returns:
            BytesIO buffer with Excel file
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Bookings"
        
        # Title and metadata
        ws.merge_cells('A1:J1')
        ws['A1'] = f"AirYatra - {title}"
        ws['A1'].font = Font(bold=True, size=16, color="FF6B00")
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws['A2'] = f"Generated: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
        ws['A2'].font = Font(italic=True, size=9, color="666666")
        
        if date_range:
            ws['A3'] = f"Period: {date_range[0]} to {date_range[1]}"
            ws['A3'].font = Font(italic=True, size=9)
            start_row = 5
        else:
            start_row = 4
        
        # Headers
        headers = [
            "Booking ID", "Date", "Customer", "Email", "Phone",
            "Route", "Passengers", "Aircraft", "Amount (₹)", "Status"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal='center')
            cell.border = THIN_BORDER
        
        # Data rows
        for row_idx, booking in enumerate(bookings, start_row + 1):
            route = f"{booking.get('from_city', 'N/A')} → {booking.get('to_city', 'N/A')}"
            
            data = [
                booking.get('id', '')[:8],
                booking.get('departure_date', booking.get('date', 'N/A')),
                booking.get('customer_name', booking.get('name', 'N/A')),
                booking.get('email', 'N/A'),
                booking.get('phone', 'N/A'),
                route,
                booking.get('passenger_count', booking.get('passengers', 1)),
                booking.get('aircraft_type', 'Helicopter'),
                booking.get('total_amount', booking.get('amount', 0)),
                booking.get('status', 'pending').title()
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal='center' if col != 6 else 'left')
                
                # Format amount column
                if col == 9:
                    cell.number_format = '₹#,##0.00'
                    cell.fill = MONEY_FILL
                
                # Color code status
                if col == 10:
                    status = str(value).lower()
                    if status == 'confirmed':
                        cell.font = Font(color="228B22")
                    elif status == 'cancelled':
                        cell.font = Font(color="DC143C")
                    elif status == 'pending':
                        cell.font = Font(color="FFA500")
        
        # Summary section
        summary_row = start_row + len(bookings) + 2
        ws.cell(row=summary_row, column=1, value="SUMMARY").font = Font(bold=True, size=12)
        
        total_bookings = len(bookings)
        total_revenue = sum(float(b.get('total_amount', b.get('amount', 0)) or 0) for b in bookings)
        confirmed = sum(1 for b in bookings if b.get('status', '').lower() == 'confirmed')
        cancelled = sum(1 for b in bookings if b.get('status', '').lower() == 'cancelled')
        total_passengers = sum(int(b.get('passenger_count', b.get('passengers', 1)) or 1) for b in bookings)
        
        summary_data = [
            ("Total Bookings:", total_bookings),
            ("Confirmed:", confirmed),
            ("Cancelled:", cancelled),
            ("Total Passengers:", total_passengers),
            ("Total Revenue:", f"₹{total_revenue:,.2f}")
        ]
        
        for idx, (label, value) in enumerate(summary_data):
            ws.cell(row=summary_row + idx + 1, column=1, value=label).font = Font(bold=True)
            ws.cell(row=summary_row + idx + 1, column=2, value=value)
        
        # Auto-adjust column widths
        self._auto_adjust_columns(ws)
        
        # Save to buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return buffer
    
    def generate_finance_report(
        self,
        transactions: List[Dict],
        title: str = "Finance Report",
        date_range: Optional[tuple] = None
    ) -> io.BytesIO:
        """Generate financial transactions report"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Transactions"
        
        # Title
        ws.merge_cells('A1:I1')
        ws['A1'] = f"AirYatra - {title}"
        ws['A1'].font = Font(bold=True, size=16, color="FF6B00")
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws['A2'] = f"Generated: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
        ws['A2'].font = Font(italic=True, size=9, color="666666")
        
        start_row = 4
        
        # Headers
        headers = [
            "Transaction ID", "Date", "Booking ID", "Customer",
            "Payment Method", "Amount (₹)", "GST (₹)", "Total (₹)", "Status"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal='center')
            cell.border = THIN_BORDER
        
        # Data rows
        for row_idx, txn in enumerate(transactions, start_row + 1):
            data = [
                txn.get('transaction_id', txn.get('id', ''))[:12],
                txn.get('created_at', 'N/A') if isinstance(txn.get('created_at'), str) else txn.get('created_at', datetime.now()).strftime('%Y-%m-%d'),
                txn.get('booking_id', 'N/A')[:8] if txn.get('booking_id') else 'N/A',
                txn.get('customer_name', txn.get('user_email', 'N/A')),
                txn.get('payment_method', 'Online').title(),
                txn.get('amount', 0) - txn.get('gst_amount', 0),
                txn.get('gst_amount', 0),
                txn.get('amount', 0),
                txn.get('status', 'success').title()
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal='center')
                
                # Format money columns
                if col in [6, 7, 8]:
                    cell.number_format = '₹#,##0.00'
                    cell.fill = MONEY_FILL
                
                # Color code status
                if col == 9:
                    status = str(value).lower()
                    if status == 'success':
                        cell.font = Font(color="228B22")
                    elif status == 'failed':
                        cell.font = Font(color="DC143C")
                    elif status == 'pending':
                        cell.font = Font(color="FFA500")
        
        # Summary section
        summary_row = start_row + len(transactions) + 2
        ws.cell(row=summary_row, column=1, value="FINANCIAL SUMMARY").font = Font(bold=True, size=12)
        
        total_amount = sum(t.get('amount', 0) or 0 for t in transactions)
        total_gst = sum(t.get('gst_amount', 0) or 0 for t in transactions)
        successful = sum(1 for t in transactions if t.get('status', '').lower() in ['success', 'completed'])
        failed = sum(1 for t in transactions if t.get('status', '').lower() == 'failed')
        
        summary_data = [
            ("Total Transactions:", len(transactions)),
            ("Successful:", successful),
            ("Failed:", failed),
            ("Gross Revenue:", f"₹{total_amount:,.2f}"),
            ("GST Collected:", f"₹{total_gst:,.2f}"),
            ("Net Revenue:", f"₹{(total_amount - total_gst):,.2f}")
        ]
        
        for idx, (label, value) in enumerate(summary_data):
            ws.cell(row=summary_row + idx + 1, column=1, value=label).font = Font(bold=True)
            ws.cell(row=summary_row + idx + 1, column=2, value=value)
        
        self._auto_adjust_columns(ws)
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return buffer
    
    def generate_refund_report(
        self,
        refunds: List[Dict],
        title: str = "Refund Report"
    ) -> io.BytesIO:
        """Generate refunds report"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Refunds"
        
        # Title
        ws.merge_cells('A1:H1')
        ws['A1'] = f"AirYatra - {title}"
        ws['A1'].font = Font(bold=True, size=16, color="FF6B00")
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws['A2'] = f"Generated: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
        
        start_row = 4
        
        headers = [
            "Refund ID", "Date", "Booking ID", "Customer",
            "Original Amount (₹)", "Refund Amount (₹)", "Reason", "Status"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER
        
        for row_idx, refund in enumerate(refunds, start_row + 1):
            data = [
                refund.get('refund_id', refund.get('id', ''))[:12],
                refund.get('created_at', 'N/A'),
                refund.get('booking_id', 'N/A')[:8],
                refund.get('customer_name', 'N/A'),
                refund.get('original_amount', 0),
                refund.get('refund_amount', 0),
                refund.get('reason', 'Customer request'),
                refund.get('status', 'pending').title()
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = THIN_BORDER
                
                if col in [5, 6]:
                    cell.number_format = '₹#,##0.00'
        
        # Summary
        summary_row = start_row + len(refunds) + 2
        total_refunded = sum(r.get('refund_amount', 0) or 0 for r in refunds)
        
        ws.cell(row=summary_row, column=1, value="Total Refunds:").font = Font(bold=True)
        ws.cell(row=summary_row, column=2, value=len(refunds))
        ws.cell(row=summary_row + 1, column=1, value="Total Amount:").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=2, value=f"₹{total_refunded:,.2f}")
        
        self._auto_adjust_columns(ws)
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return buffer
    
    def generate_customer_report(
        self,
        customers: List[Dict],
        title: str = "Customer Report"
    ) -> io.BytesIO:
        """Generate customer analytics report"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Customers"
        
        ws.merge_cells('A1:H1')
        ws['A1'] = f"AirYatra - {title}"
        ws['A1'].font = Font(bold=True, size=16, color="FF6B00")
        ws['A1'].alignment = Alignment(horizontal='center')
        
        start_row = 4
        
        headers = [
            "Customer ID", "Name", "Email", "Phone",
            "Total Bookings", "Total Spent (₹)", "Loyalty Points", "Status"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER
        
        for row_idx, customer in enumerate(customers, start_row + 1):
            data = [
                customer.get('id', '')[:8],
                customer.get('full_name', customer.get('name', 'N/A')),
                customer.get('email', 'N/A'),
                customer.get('phone', 'N/A'),
                customer.get('total_bookings', 0),
                customer.get('total_spent', 0),
                customer.get('loyalty_points', 0),
                'Active' if customer.get('is_active', True) else 'Inactive'
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = THIN_BORDER
                
                if col == 6:
                    cell.number_format = '₹#,##0.00'
                    cell.fill = MONEY_FILL
        
        self._auto_adjust_columns(ws)
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return buffer
    
    def _auto_adjust_columns(self, ws):
        """Auto-adjust column widths based on content"""
        for column_cells in ws.columns:
            max_length = 0
            column = None
            
            for cell in column_cells:
                try:
                    # Skip merged cells
                    if hasattr(cell, 'column_letter'):
                        if column is None:
                            column = cell.column_letter
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            
            if column:
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column].width = adjusted_width


# Singleton instance
report_export_service = ReportExportService()
