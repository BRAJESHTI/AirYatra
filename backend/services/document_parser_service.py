"""
AirYatra Bulk Document Import Service
Parses Word (.docx) and PDF files for importing legal documents
"""

import io
import logging
import re
from typing import Dict, Any, Optional, List
from datetime import datetime

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

logger = logging.getLogger(__name__)


class DocumentParser:
    """Parser for Word and PDF documents"""
    
    def __init__(self):
        self.supported_types = []
        if DOCX_AVAILABLE:
            self.supported_types.extend(['.docx', '.doc'])
        if PDF_AVAILABLE:
            self.supported_types.extend(['.pdf'])
    
    def get_supported_types(self) -> List[str]:
        """Get list of supported file extensions"""
        return self.supported_types
    
    def parse_file(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """
        Parse uploaded file and extract content
        
        Args:
            file_content: Raw file bytes
            filename: Original filename with extension
            
        Returns:
            Dict with success, content (plain text), html_content, metadata
        """
        ext = '.' + filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
        
        if ext in ['.docx', '.doc']:
            return self._parse_docx(file_content, filename)
        elif ext == '.pdf':
            return self._parse_pdf(file_content, filename)
        else:
            return {
                "success": False,
                "error": f"Unsupported file type: {ext}. Supported: {', '.join(self.supported_types)}"
            }
    
    def _parse_docx(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Parse Word document"""
        if not DOCX_AVAILABLE:
            return {"success": False, "error": "python-docx library not installed"}
        
        try:
            doc = DocxDocument(io.BytesIO(file_content))
            
            # Extract plain text
            plain_text = []
            html_parts = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                
                plain_text.append(text)
                
                # Detect headings by style
                style_name = para.style.name.lower() if para.style else ''
                
                if 'heading 1' in style_name or 'title' in style_name:
                    html_parts.append(f"<h1>{self._escape_html(text)}</h1>")
                elif 'heading 2' in style_name:
                    html_parts.append(f"<h2>{self._escape_html(text)}</h2>")
                elif 'heading 3' in style_name:
                    html_parts.append(f"<h3>{self._escape_html(text)}</h3>")
                elif text.startswith('•') or text.startswith('-') or text.startswith('*'):
                    html_parts.append(f"<li>{self._escape_html(text[1:].strip())}</li>")
                else:
                    html_parts.append(f"<p>{self._escape_html(text)}</p>")
            
            # Wrap list items
            html_content = '\n'.join(html_parts)
            html_content = self._wrap_list_items(html_content)
            
            # Extract metadata
            core_props = doc.core_properties
            metadata = {
                "author": core_props.author,
                "title": core_props.title,
                "created": core_props.created.isoformat() if core_props.created else None,
                "modified": core_props.modified.isoformat() if core_props.modified else None,
                "paragraph_count": len(doc.paragraphs),
                "word_count": len(' '.join(plain_text).split())
            }
            
            # Try to detect document type from title/content
            detected_type = self._detect_document_type('\n'.join(plain_text))
            
            return {
                "success": True,
                "filename": filename,
                "file_type": "docx",
                "plain_text": '\n\n'.join(plain_text),
                "html_content": html_content,
                "metadata": metadata,
                "detected_doc_type": detected_type,
                "suggested_title": metadata.get("title") or self._extract_title(plain_text)
            }
            
        except Exception as e:
            logger.error(f"Error parsing DOCX: {e}")
            return {"success": False, "error": f"Failed to parse Word document: {str(e)}"}
    
    def _parse_pdf(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Parse PDF document"""
        if not PDF_AVAILABLE:
            return {"success": False, "error": "PyPDF2 library not installed"}
        
        try:
            reader = PyPDF2.PdfReader(io.BytesIO(file_content))
            
            plain_text = []
            html_parts = []
            
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    # Clean up extracted text
                    lines = text.split('\n')
                    for line in lines:
                        line = line.strip()
                        if not line:
                            continue
                        
                        plain_text.append(line)
                        
                        # Simple heuristic for headings (all caps or short bold-like text)
                        if line.isupper() and len(line) < 100:
                            html_parts.append(f"<h2>{self._escape_html(line.title())}</h2>")
                        elif len(line) < 60 and not line.endswith('.'):
                            html_parts.append(f"<h3>{self._escape_html(line)}</h3>")
                        else:
                            html_parts.append(f"<p>{self._escape_html(line)}</p>")
            
            html_content = '\n'.join(html_parts)
            
            # Extract metadata
            metadata = {
                "page_count": len(reader.pages),
                "word_count": len(' '.join(plain_text).split()),
                "pdf_info": {}
            }
            
            if reader.metadata:
                metadata["pdf_info"] = {
                    "title": reader.metadata.get('/Title'),
                    "author": reader.metadata.get('/Author'),
                    "creator": reader.metadata.get('/Creator'),
                    "created": reader.metadata.get('/CreationDate'),
                }
            
            # Detect document type
            detected_type = self._detect_document_type('\n'.join(plain_text))
            
            return {
                "success": True,
                "filename": filename,
                "file_type": "pdf",
                "plain_text": '\n\n'.join(plain_text),
                "html_content": html_content,
                "metadata": metadata,
                "detected_doc_type": detected_type,
                "suggested_title": metadata.get("pdf_info", {}).get("title") or self._extract_title(plain_text)
            }
            
        except Exception as e:
            logger.error(f"Error parsing PDF: {e}")
            return {"success": False, "error": f"Failed to parse PDF: {str(e)}"}
    
    def _escape_html(self, text: str) -> str:
        """Escape HTML special characters"""
        return (text
            .replace('&', '&amp;')
            .replace('<', '&lt;')
            .replace('>', '&gt;')
            .replace('"', '&quot;')
        )
    
    def _wrap_list_items(self, html: str) -> str:
        """Wrap consecutive <li> elements in <ul>"""
        # Simple approach: wrap any <li> sequences
        html = re.sub(r'(<li>.*?</li>\n?)+', r'<ul>\g<0></ul>', html, flags=re.DOTALL)
        return html
    
    def _detect_document_type(self, text: str) -> Optional[str]:
        """Detect legal document type from content"""
        text_lower = text.lower()
        
        type_keywords = {
            "terms_conditions": ["terms and conditions", "terms of service", "terms of use", "user agreement"],
            "privacy_policy": ["privacy policy", "privacy notice", "data protection", "personal information"],
            "refund_policy": ["refund policy", "cancellation policy", "refund and cancellation", "return policy"],
            "operator_agreement": ["operator agreement", "service provider agreement", "partnership agreement"],
            "booking_terms": ["booking terms", "reservation terms", "booking conditions"],
            "safety_guidelines": ["safety guidelines", "safety instructions", "safety policy", "safety rules"],
            "cookie_policy": ["cookie policy", "cookie notice", "use of cookies"]
        }
        
        for doc_type, keywords in type_keywords.items():
            for keyword in keywords:
                if keyword in text_lower:
                    return doc_type
        
        return None
    
    def _extract_title(self, text_lines: List[str]) -> Optional[str]:
        """Extract suggested title from first meaningful line"""
        for line in text_lines[:5]:  # Check first 5 lines
            line = line.strip()
            if line and len(line) < 100 and not line.endswith('.'):
                return line
        return None


# Singleton instance
document_parser = DocumentParser()
