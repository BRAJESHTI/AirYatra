"""
Flight Log Export Service
Export pilot flight logs to PDF and Excel formats
"""
import io
from datetime import datetime
from typing import List, Dict, Any

# PDF Generation
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

# Excel Generation
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


class FlightLogExporter:
    """Export flight logs to various formats"""
    
    def __init__(self):
        self.styles = getSampleStyleSheet()
        self._setup_custom_styles()
    
    def _setup_custom_styles(self):
        """Setup custom paragraph styles for PDF"""
        self.styles.add(ParagraphStyle(
            name='Title_Custom',
            parent=self.styles['Heading1'],
            fontSize=20,
            textColor=colors.HexColor('#f97316'),
            alignment=TA_CENTER,
            spaceAfter=10
        ))
        
        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=11,
            textColor=colors.HexColor('#64748b'),
            alignment=TA_CENTER,
            spaceAfter=20
        ))
    
    def generate_pdf(self, pilot_info: Dict, flight_logs: List[Dict]) -> bytes:
        """
        Generate PDF flight log book
        
        Args:
            pilot_info: Dictionary with pilot details (name, license, email)
            flight_logs: List of flight log dictionaries
            
        Returns:
            PDF file as bytes
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=15*mm,
            leftMargin=15*mm,
            topMargin=15*mm,
            bottomMargin=15*mm
        )
        
        elements = []
        
        # Header
        elements.append(Paragraph("AirYatra Pilot Logbook", self.styles['Title_Custom']))
        elements.append(Paragraph(
            f"Flight Log Export - {pilot_info.get('name', 'Pilot')}",
            self.styles['Subtitle']
        ))
        
        # Pilot Info
        pilot_info_text = f"""
        <b>Pilot:</b> {pilot_info.get('name', 'N/A')} | 
        <b>License:</b> {pilot_info.get('license_number', 'N/A')} | 
        <b>Email:</b> {pilot_info.get('email', 'N/A')}
        """
        elements.append(Paragraph(pilot_info_text, ParagraphStyle(
            name='PilotInfo',
            fontSize=10,
            textColor=colors.HexColor('#475569'),
            alignment=TA_CENTER,
            spaceAfter=15
        )))
        
        # Summary Stats
        total_flights = len(flight_logs)
        total_hours = sum(log.get('duration_hours', log.get('duration', 0)) for log in flight_logs)
        
        summary_data = [
            ['Total Flights', 'Total Hours', 'Export Date'],
            [str(total_flights), f"{total_hours:.1f}h", datetime.now().strftime('%d %B %Y')]
        ]
        
        summary_table = Table(summary_data, colWidths=[80, 80, 120])
        summary_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#64748b')),
            ('TEXTCOLOR', (0, 1), (-1, 1), colors.HexColor('#1e293b')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 15))
        
        # Flight Log Table Header
        table_header = [
            'S.No', 'Date', 'Flight ID', 'From', 'To', 'Aircraft', 
            'Duration', 'Type', 'PIC/SIC', 'Remarks'
        ]
        
        # Flight Log Data
        table_data = [table_header]
        for idx, log in enumerate(flight_logs, 1):
            row = [
                str(idx),
                log.get('date', 'N/A'),
                log.get('flight_id', 'N/A'),
                log.get('departure_location', log.get('from', 'N/A')),
                log.get('arrival_location', log.get('to', 'N/A')),
                log.get('aircraft_registration', log.get('aircraft', 'N/A')),
                f"{log.get('duration_hours', log.get('duration', 0)):.1f}h",
                log.get('flight_type', 'N/A'),
                log.get('role', 'PIC'),
                log.get('remarks', '')[:20]
            ]
            table_data.append(row)
        
        # Create table
        col_widths = [30, 70, 80, 70, 70, 70, 50, 70, 50, 100]
        flight_table = Table(table_data, colWidths=col_widths)
        
        flight_table.setStyle(TableStyle([
            # Header
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f97316')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            
            # Data rows
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ALIGN', (0, 1), (0, -1), 'CENTER'),  # S.No
            ('ALIGN', (6, 1), (6, -1), 'CENTER'),  # Duration
            
            # Grid
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            
            # Padding
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
        ]))
        
        elements.append(flight_table)
        elements.append(Spacer(1, 20))
        
        # Footer
        elements.append(Paragraph(
            f"Generated by AirYatra Aviation Platform | {datetime.now().strftime('%d %B %Y, %H:%M IST')}",
            ParagraphStyle(
                name='Footer',
                fontSize=8,
                textColor=colors.HexColor('#94a3b8'),
                alignment=TA_CENTER
            )
        ))
        
        elements.append(Paragraph(
            "This document is a computer-generated flight log export and is for record purposes only.",
            ParagraphStyle(
                name='Disclaimer',
                fontSize=7,
                textColor=colors.HexColor('#94a3b8'),
                alignment=TA_CENTER
            )
        ))
        
        # Build PDF
        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()
    
    def generate_excel(self, pilot_info: Dict, flight_logs: List[Dict]) -> bytes:
        """
        Generate Excel flight log book
        
        Args:
            pilot_info: Dictionary with pilot details
            flight_logs: List of flight log dictionaries
            
        Returns:
            Excel file as bytes
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Flight Logs"
        
        # Styles
        header_fill = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=14, color="F97316")
        border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        
        # Title
        ws.merge_cells('A1:J1')
        ws['A1'] = "AirYatra Pilot Flight Logbook"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Pilot Info
        ws.merge_cells('A2:J2')
        ws['A2'] = f"Pilot: {pilot_info.get('name', 'N/A')} | License: {pilot_info.get('license_number', 'N/A')} | Email: {pilot_info.get('email', 'N/A')}"
        ws['A2'].alignment = Alignment(horizontal='center')
        
        # Summary
        total_flights = len(flight_logs)
        total_hours = sum(log.get('duration_hours', log.get('duration', 0)) for log in flight_logs)
        
        ws['A4'] = "Total Flights:"
        ws['B4'] = total_flights
        ws['C4'] = "Total Hours:"
        ws['D4'] = f"{total_hours:.1f}h"
        ws['E4'] = "Export Date:"
        ws['F4'] = datetime.now().strftime('%d %B %Y')
        
        # Headers (Row 6)
        headers = ['S.No', 'Date', 'Flight ID', 'From', 'To', 'Aircraft', 
                   'Duration (h)', 'Type', 'Role', 'Remarks']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        
        # Data rows
        for idx, log in enumerate(flight_logs, 1):
            row = idx + 6
            data = [
                idx,
                log.get('date', 'N/A'),
                log.get('flight_id', 'N/A'),
                log.get('departure_location', log.get('from', 'N/A')),
                log.get('arrival_location', log.get('to', 'N/A')),
                log.get('aircraft_registration', log.get('aircraft', 'N/A')),
                round(log.get('duration_hours', log.get('duration', 0)), 1),
                log.get('flight_type', 'N/A'),
                log.get('role', 'PIC'),
                log.get('remarks', '')
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = border
                cell.alignment = Alignment(horizontal='center' if col in [1, 7] else 'left')
        
        # Column widths
        col_widths = [6, 12, 15, 15, 15, 12, 12, 12, 8, 25]
        for idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        
        # Freeze header row
        ws.freeze_panes = 'A7'
        
        # Save to buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


# Singleton instance
flight_log_exporter = FlightLogExporter()
